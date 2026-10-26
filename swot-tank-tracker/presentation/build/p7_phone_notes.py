"""Practice notes for a phone: each slide's picture with its speaker notes, one A5 page per slide.

Needs the slide pictures from render.ps1 (render/slide-NN.png). Prints with Microsoft Edge.
Usage: python p7_phone_notes.py <output.pdf>
"""
import base64
import html
import subprocess
import sys
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

BUILD = Path(__file__).resolve().parent
ROOT = BUILD.parent
DECK = ROOT / "SWOT Tank Water Tracker - proposal talk.pptx"
RENDER = ROOT / "render"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
WORDS_PER_MINUTE = 130  # a calm speaking pace

CSS = """
@page { size: A5; margin: 10mm 10mm 13mm;
  @bottom-center { content: counter(page) " / " counter(pages); font: 8pt "Segoe UI"; color: #56656F; } }
html { font-family: "Segoe UI", sans-serif; font-size: 11pt; color: #142430; background: #fff; }
body { margin: 0; }
section { break-after: page; }
section:last-child { break-after: auto; }
.top { font-size: 8.5pt; color: #56656F; margin-bottom: 2mm; }
h1 { font-family: "Segoe UI Semibold", "Segoe UI", sans-serif; font-weight: 600; font-size: 13pt;
  line-height: 1.25; margin: 0 0 3mm; color: #0B2230; }
h1 .n { color: #2A78D6; margin-right: 1.5mm; }
img { display: block; width: 100%; border: 0.3mm solid #C9D3DA; border-radius: 1.5mm; margin-bottom: 4mm; }
p { margin: 0 0 2.5mm; line-height: 1.5; }
"""


def title_of(slide):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.TITLE:
            return sh.text_frame.text.strip()
    return ""


def picture(path):
    img = Image.open(path).convert("RGB")
    img.thumbnail((1200, 1200))
    buf = BytesIO()
    img.save(buf, "JPEG", quality=85)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def main(out_pdf):
    slides = list(Presentation(DECK).slides)
    pngs = sorted(RENDER.glob("slide-*.png"))
    assert len(pngs) == len(slides), f"{len(pngs)} pictures for {len(slides)} slides; run render.ps1 first"
    assert all(p.stat().st_mtime >= DECK.stat().st_mtime for p in pngs), "slide pictures are older than the deck"

    notes = [s.notes_slide.notes_text_frame.text.replace("\v", " ") if s.has_notes_slide else "" for s in slides]
    minutes = round(sum(len(n.split()) for n in notes) / WORDS_PER_MINUTE)
    pages = []
    for i, (slide, png, note) in enumerate(zip(slides, pngs, notes), start=1):
        top = (f'<div class="top">Speaker notes · {len(slides)} slides · about {minutes} minutes at a calm pace</div>'
               if i == 1 else "")
        paras = "".join(f"<p>{html.escape(p.strip())}</p>" for p in note.split("\n") if p.strip())
        pages.append(f'<section>{top}<h1><span class="n">{i}</span>{html.escape(title_of(slide))}</h1>'
                     f'<img src="{picture(png)}" alt="Slide {i}">{paras}</section>')

    page = RENDER / "phone_notes.html"
    page.write_text(f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Speaker notes</title>'
                    f"<style>{CSS}</style></head><body>{''.join(pages)}</body></html>", encoding="utf-8")
    out_pdf = Path(out_pdf).resolve()
    if out_pdf.exists():
        out_pdf.unlink()
    subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--no-first-run", "--no-pdf-header-footer",
                    f"--user-data-dir={BUILD / 'edge-profile'}", f"--print-to-pdf={out_pdf}",
                    "--virtual-time-budget=10000", page.as_uri()], check=True, timeout=240, capture_output=True)
    assert out_pdf.exists(), "Edge did not write the PDF"
    print(f"{out_pdf.name}: {len(slides)} slides, about {minutes} minutes of notes")


if __name__ == "__main__":
    main(sys.argv[1])
